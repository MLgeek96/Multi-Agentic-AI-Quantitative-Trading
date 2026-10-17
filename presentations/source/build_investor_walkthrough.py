"""Build the investor speaking script and matching plain-text rehearsal copy.

Run with the Python executable returned by Codex's workspace dependency loader.
The document is tied to the implementation reviewed on 17 October 2026.
"""

from pathlib import Path
import json
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / "presentations" / "roadshow_20261017"
DOCX = OUT / "Investor_System_Walkthrough_and_QA.docx"
TEXT = OUT / "Investor_System_Walkthrough_and_QA.txt"

from pairs_walkthrough_content import CHAPTERS, QA, PITCH, PREP, CLAIMS, SOURCE_NOTES


def words(text):
    return len(re.findall(r"\S+", text))


def set_font(style, size, bold=False):
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = OxmlElement("w:rFonts")
        rpr.append(fonts)
    for name in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn("w:" + name), "Calibri")
    color = rpr.find(qn("w:color"))
    if color is not None:
        for name in ("themeColor", "themeTint", "themeShade"):
            color.attrib.pop(qn("w:" + name), None)


def add_p(doc, text, *, style=None, bold_lead=None):
    p = doc.add_paragraph(style=style)
    if bold_lead:
        p.add_run(bold_lead + " ").bold = True
    p.add_run(text)
    return p


def table(doc, headings, rows, widths):
    t = doc.add_table(rows=1, cols=len(headings))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    for c, width in zip(t.columns, widths):
        c.width = Inches(width)
    for i, h in enumerate(headings):
        t.rows[0].cells[i].text = h
    for row in rows:
        cells = t.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = text
    for ri, row in enumerate(t.rows):
        props = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        props.append(cant)
        if ri == 0:
            props.append(OxmlElement("w:tblHeader"))
        for ci, cell in enumerate(row.cells):
            cell.width = Inches(widths[ci])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement("w:tcMar")
            for edge in ("top", "left", "bottom", "right"):
                e = OxmlElement("w:" + edge)
                e.set(qn("w:w"), "105")
                e.set(qn("w:type"), "dxa")
                margins.append(e)
            tcpr.append(margins)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                e = OxmlElement("w:" + edge)
                e.set(qn("w:val"), "single")
                e.set(qn("w:sz"), "4")
                e.set(qn("w:color"), "D9D9D9")
                borders.append(e)
            tcpr.append(borders)
            shade = OxmlElement("w:shd")
            shade.set(qn("w:fill"), "E7EEF2" if ri == 0 else ("F5F7F9" if ri % 2 == 0 else "FFFFFF"))
            tcpr.append(shade)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.04
                for run in p.runs:
                    run.font.size = Pt(10)
                    run.font.color.rgb = RGBColor(0, 0, 0)
                    run.bold = ri == 0
    add_p(doc, "").paragraph_format.space_after = Pt(2)
    return t


def build():
    OUT.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.top_margin = sec.bottom_margin = Inches(0.7)
    sec.left_margin = sec.right_margin = Inches(0.78)
    for name, size, bold in [("Normal", 11.5, False), ("Title", 26, True), ("Subtitle", 12, False), ("Heading 1", 17, True), ("Heading 2", 13, True), ("Heading 3", 11.5, True)]:
        set_font(doc.styles[name], size, bold)
    # The bundled starter has an inherited blue bottom border in its Title style.
    # Remove paragraph borders so titles and headings use typography alone.
    for style in doc.styles:
        for border in list(style.element.iter(qn("w:pBdr"))):
            border.getparent().remove(border)
    normal = doc.styles["Normal"].paragraph_format
    normal.space_after = Pt(8)
    normal.line_spacing = 1.12
    normal.widow_control = True
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        fmt = doc.styles[name].paragraph_format
        fmt.space_before, fmt.space_after = Pt(13), Pt(7)
        fmt.keep_with_next = True
    doc.styles["Title"].paragraph_format.space_after = Pt(10)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Page ")
    r.font.size = Pt(9)
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    footer._p.append(fld)
    doc.core_properties.title = "Investor system walkthrough and questions"
    doc.core_properties.subject = "FE5209 quantitative paper fund speaking script"
    doc.core_properties.author = "FE5209 project team"
    doc.core_properties.keywords = "investor walkthrough, quantitative fund, presentation script"

    speech_words = sum(words(p) for ch in CHAPTERS for p in ch["say"])
    minutes_low, minutes_high = round(speech_words / 150 + 2), round(speech_words / 130 + 3)
    add_p(doc, "Investor system walkthrough and questions", style="Title")
    add_p(doc, "FE5209 quantitative paper fund 17 October 2026", style="Subtitle")
    add_p(doc, "A complete speaking script for explaining the fund, demonstrating one saved decision, and answering investor questions. The central proposition is an inspectable decision process: quantitative engines propose, independent risk decides, and AI reviews the evidence.")
    add_p(doc, f"Main talk: {speech_words:,} spoken words. Allow approximately {minutes_low} to {minutes_high} minutes including brief screen changes, before questions. Timing is a rehearsal estimate, not a measured recording.", bold_lead="Rehearsal length")
    table(doc, ["Part", "Use"], [
        ("Main walkthrough", "Twelve speaking chapters with screen cues and suggested presenters"),
        ("Investor questions", "Fifteen answers about strategy, AI, risk, evidence and readiness"),
        ("One minute pitch", "A shorter narration for a 60 second recording"),
        ("Presenter preparation", "Demo setup, fallback files, claim boundaries and evidence references"),
    ], [1.5, 5.44])
    add_p(doc, "The current integrated prototype connects only Pairs. Crypto is removed from the active fund. Momentum, automatic model calls, fills, and consolidated fund performance remain next milestones. Keep those boundaries explicit when presenting the system.", bold_lead="Current scope")

    doc.add_page_break()
    add_p(doc, "Main investor walkthrough", style="Heading 1")
    plain = ["INVESTOR SYSTEM WALKTHROUGH AND QUESTIONS", "FE5209 quantitative paper fund | 17 October 2026", "", f"MAIN TALK: {speech_words:,} spoken words; estimated {minutes_low}-{minutes_high} minutes including screen changes. Rehearse to confirm.", "Read SAY paragraphs aloud. SHOW and PRESENTER are rehearsal cues.", ""]
    for ch in CHAPTERS:
        add_p(doc, ch["title"], style="Heading 2")
        p = add_p(doc, ch["screen"], bold_lead="SHOW")
        p.paragraph_format.keep_with_next = True
        for run in p.runs:
            run.font.size = Pt(10.5)
        p = add_p(doc, ch["speaker"], bold_lead="PRESENTER")
        p.paragraph_format.keep_with_next = True
        for run in p.runs:
            run.font.size = Pt(10.5)
        for i, para in enumerate(ch["say"]):
            add_p(doc, para, bold_lead="SAY" if i == 0 else None)
        plain.extend([ch["title"].upper(), "SHOW: " + ch["screen"], "PRESENTER: " + ch["speaker"], "", "SAY:", *[p + "\n" for p in ch["say"]], ""])

    doc.add_page_break()
    add_p(doc, "Investor questions and suggested answers", style="Heading 1")
    add_p(doc, "Use the answer that addresses the question directly, then offer to show the relevant saved evidence. Do not treat proposed experiments as completed work.")
    plain.extend(["INVESTOR QUESTIONS AND SUGGESTED ANSWERS", ""])
    for i, (question, answer) in enumerate(QA, 1):
        p = add_p(doc, f"{i} {question}", style="Heading 2")
        p.paragraph_format.space_before = Pt(9)
        p.paragraph_format.space_after = Pt(5)
        p = add_p(doc, answer)
        p.paragraph_format.keep_together = True
        p.paragraph_format.line_spacing = 1.08
        p.paragraph_format.space_after = Pt(6)
        plain.extend([f"{i}. {question}?", answer, ""])

    doc.add_page_break()
    add_p(doc, "One minute investor pitch", style="Heading 1")
    add_p(doc, f"{words(PITCH)} whitespace-delimited words. Rehearse for 55 to 60 seconds; actual delivery determines the duration.")
    add_p(doc, PITCH)
    add_p(doc, "Recording sequence", style="Heading 2")
    table(doc, ["Approximate time", "Visual"], [
        ("0 to 10 seconds", "Fund overview and proposed paper mandate"),
        ("10 to 25 seconds", "System flow and independent fund risk"),
        ("25 to 45 seconds", "Linked targets, diagnostics and the saved approval"),
        ("45 to 60 seconds", "Zero fills label and next paper-pilot milestone"),
    ], [1.5, 5.44])
    plain.extend(["ONE MINUTE INVESTOR PITCH", f"{words(PITCH)} words; rehearse duration.", "", PITCH, ""])

    doc.add_page_break()
    add_p(doc, "Presenter preparation and evidence", style="Heading 1")
    plain.extend(["PRESENTER PREPARATION AND EVIDENCE", ""])
    for title, paras in PREP:
        add_p(doc, title, style="Heading 2")
        plain.extend([title.upper(), ""])
        for para in paras:
            add_p(doc, para)
            plain.extend([para, ""])
    add_p(doc, "Claims to keep precise", style="Heading 2")
    table(doc, ["Subject", "Supported wording", "Boundary"], CLAIMS, [1.42, 2.72, 2.8])
    plain.extend(["CLAIMS TO KEEP PRECISE", ""])
    for subject, claim, boundary in CLAIMS:
        plain.extend([subject + ": " + claim, "Boundary: " + boundary, ""])
    doc.add_page_break()
    add_p(doc, "Evidence references", style="Heading 1")
    plain.extend(["EVIDENCE REFERENCES", ""])
    for i, note in enumerate(SOURCE_NOTES, 1):
        add_p(doc, f"[{i}] {note}")
        plain.extend([f"[{i}] {note}", ""])
    doc.save(DOCX)
    TEXT.write_text("\n".join(plain), encoding="utf-8")
    stats = {"docx": str(DOCX), "text": str(TEXT), "chapters": len(CHAPTERS), "main_spoken_words": speech_words, "pitch_words": words(PITCH), "estimated_minutes": [minutes_low, minutes_high], "questions": len(QA)}
    print(json.dumps(stats, indent=2))


if __name__ == "__main__":
    build()
